"""Fill the ECA Group Assessment Cover Sheet for one assessment.

Reads the untouched template (templates/Group_Assessment_Cover_Sheet.docx), the shared
reports/course_details.json and reports/<A>/cover_sheet.json, and writes a filled copy.
Unknown values stay as yellow-highlighted [[placeholders]]. Declaration checkboxes are only
ticked when the students have set them to true in cover_sheet.json.

Usage: python reports/tools/fill_cover_sheet.py A1
"""
import copy
import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml.ns import qn
from docx.shared import Pt

sys.path.insert(0, str(Path(__file__).parent))
from wordcount import count  # noqa: E402

ROOT = Path(__file__).resolve().parents[2]
TEMPLATE = ROOT / "templates" / "Group_Assessment_Cover_Sheet.docx"
TEMPLATE_INDIVIDUAL = ROOT / "templates" / "Individual_Assessment_Cover_Sheet.docx"
W14 = "http://schemas.microsoft.com/office/word/2010/wordml"
FONT, SIZE = "Arial", Pt(10.5)


def write_cell(cell, text):
    p = cell.paragraphs[0]
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    run = p.add_run(text)
    run.font.name, run.font.size = FONT, SIZE
    if "[[" in text:
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def set_checkbox(sdt, checked):
    state = sdt.find(f".//{{{W14}}}checked")
    state.set(f"{{{W14}}}val", "1" if checked else "0")
    for t in sdt.iter(qn("w:t")):
        t.text = "☒" if checked else "☐"


def replace_paragraph_text(paragraph, text):
    for r in list(paragraph.runs):
        r._r.getparent().remove(r._r)
    run = paragraph.add_run(text)
    run.font.name, run.font.size = FONT, SIZE


def fill(assessment: str) -> Path:
    course = json.loads((ROOT / "reports" / "course_details.json").read_text(encoding="utf-8"))
    folder = ROOT / "reports" / assessment
    cfg = json.loads((folder / "cover_sheet.json").read_text(encoding="utf-8"))

    word_count = cfg.get("word_count", "auto")  # "auto" = counted from the report; "" = left blank
    if word_count == "auto":
        word_count = f"{count(str(folder / cfg['report_markdown']))['prose']} words (body text)"

    values = [
        course["course_title"], course["unit"], course["lecturer"], course["trimester"],
        cfg["assessment_number"], cfg["assessment_title"], word_count, cfg["due_date"],
    ]
    if cfg.get("template") == "individual":
        # One table: name, ID, the eight shared fields, then the three declaration rows.
        doc = Document(str(TEMPLATE_INDIVIDUAL))
        details = doc.tables[0]
        student = course["students"][cfg["student_index"]]
        write_cell(details.rows[0].cells[1], student["name"])
        write_cell(details.rows[1].cells[1], student["id"])
        for row, value in zip(details.rows[2:10], values):
            write_cell(row.cells[1], value)
        decl_rows = details.rows[10:13]
    else:
        doc = Document(str(TEMPLATE))
        students, details = doc.tables[0], doc.tables[1]
        for row, student in zip(students.rows[1:], course["students"]):
            write_cell(row.cells[0], student["name"])
            write_cell(row.cells[1], student["id"])
        for row, value in zip(details.rows[:8], values):
            write_cell(row.cells[1], value)
        decl_rows = details.rows[8:11]

    decl = cfg.get("declarations", {})
    set_checkbox(decl_rows[0].cells[1]._tc.find(f".//{qn('w:sdt')}"), bool(decl.get("student")))
    set_checkbox(decl_rows[1].cells[1]._tc.find(f".//{qn('w:sdt')}"), bool(decl.get("reference")))
    ai_cell = decl_rows[2].cells[1]
    none_box, used_box = list(ai_cell._tc.iter(qn("w:sdt")))[:2]
    set_checkbox(none_box, decl.get("ai_choice") == "none")
    set_checkbox(used_box, decl.get("ai_choice") == "used")

    # Template paragraphs 5-7 hold "We acknowledge the use of [tool] ... to [how it was used]";
    # 8-9 are blank spacers before "Please refer to". Keep the sheet on one page.
    acks = cfg.get("ai_acknowledgements", [])
    paras = ai_cell.paragraphs
    first, filler = paras[5], paras[6:9]
    if acks:
        replace_paragraph_text(first, acks[0])
        fmt = first.paragraph_format
        fmt.line_spacing, fmt.space_before, fmt.space_after = 1.0, Pt(2), Pt(2)
        anchor = first._p
        for text in acks[1:]:
            new_p = copy.deepcopy(first._p)
            anchor.addnext(new_p)
            anchor = new_p
            for t in new_p.iter(qn("w:t")):
                t.text = ""
            new_p.findall(qn("w:r"))[0].find(qn("w:t")).text = text
        for p in filler:
            p._p.getparent().remove(p._p)

    out_dir = folder / "submission"
    out_dir.mkdir(exist_ok=True)
    out = out_dir / f"{Path(assessment).name}_Cover_Sheet.docx"
    doc.save(str(out))
    print(f"Wrote {out}")
    return out


if __name__ == "__main__":
    fill(sys.argv[1])
