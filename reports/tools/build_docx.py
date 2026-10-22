"""Build a submission-ready Word report from a Markdown draft.

Pipeline: Markdown -> pandoc (with a styled reference doc) -> python-docx post-processing
(cover page, APA tables, figure layout, hanging-indent references, page numbers,
yellow-highlighted [[placeholders]]).

Usage:
    python reports/tools/build_docx.py reports/A1/A1_Project_Brief.md reports/A1/A1_Project_Brief.docx
"""
import copy
import re
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph
from docx.text.run import Run

sys.path.insert(0, str(Path(__file__).parent))
from wordcount import CAPTION_LABEL, count  # noqa: E402

FONT = "Calibri"
BODY_PT = 11
NAVY = RGBColor(0x1F, 0x38, 0x64)
TABLE_SHADE = "E8EEF7"


def apply_theme(meta):
    """Per-report look from front matter: font, body_size, heading_color (hex), table_shade (hex)."""
    global FONT, BODY_PT, NAVY, TABLE_SHADE
    FONT = meta.get("font", "Calibri")
    BODY_PT = float(meta.get("body_size", 11))
    NAVY = RGBColor.from_string(meta.get("heading_color", "1F3864").lstrip("#"))
    TABLE_SHADE = meta.get("table_shade", "E8EEF7").lstrip("#")


# ---------- front matter ----------
def split_front_matter(md: str) -> tuple[dict, str]:
    m = re.match(r"^---\n(.*?)\n---\n", md, flags=re.S)
    if not m:
        return {}, md
    meta, key = {}, None
    for line in m.group(1).splitlines():
        if re.match(r"^\s+-\s", line) and key:
            meta.setdefault(key, []).append(line.split("-", 1)[1].strip().strip('"'))
        elif ":" in line:
            key, val = line.split(":", 1)
            key, val = key.strip(), val.strip().strip('"')
            meta[key] = val if val else []
    return meta, md[m.end():]


def prepare_markdown(body: str) -> str:
    lines = body.splitlines()
    out = []
    for i, line in enumerate(lines):
        # APA caption: bold label line followed by italic title line -> keep on separate lines
        if CAPTION_LABEL.fullmatch(line.strip()) and i + 1 < len(lines) and lines[i + 1].strip():
            line = line.rstrip() + "\\"
        line = re.sub(r"!\[\]\((.+?)\)$", r"![](\1){width=16cm}", line)
        out.append(line)
    return "\n".join(out) + "\n"


# ---------- styles ----------
def set_style_font(style, size=None, bold=None, italic=None, color=None, name=None):
    name = name or FONT
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        fonts.attrib.pop(qn(attr), None)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), name)
    if size is not None:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic
    if color is not None:
        style.font.color.rgb = color


def make_reference_doc(path: Path) -> None:
    data = subprocess.run(
        ["pandoc", "-o", str(path), "--print-default-data-file", "reference.docx"],
        check=True,
    )
    doc = Document(str(path))
    st = doc.styles
    for name in ("Normal", "Body Text", "First Paragraph", "Compact", "Block Text"):
        if name in [s.name for s in st]:
            s = st[name]
            set_style_font(s, size=BODY_PT, color=RGBColor(0, 0, 0))
            pf = s.paragraph_format
            pf.line_spacing = 1.5 if name != "Compact" else 1.3
            pf.space_before = Pt(0)
            pf.space_after = Pt(6 if name != "Compact" else 3)
    for name, size in (("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 11)):
        s = st[name]
        set_style_font(s, size=size, bold=True, italic=False, color=NAVY)
        s.paragraph_format.space_before = Pt(14 if name == "Heading 1" else 10)
        s.paragraph_format.space_after = Pt(6)
        s.paragraph_format.keep_with_next = True
    if "Hyperlink" in [s.name for s in st]:
        set_style_font(st["Hyperlink"], color=RGBColor(0x05, 0x63, 0xC1))
    set_language(doc, "en-AU")
    doc.save(str(path))


def set_language(doc, lang):
    """Proofing language for the whole document (Australian English spelling check)."""
    styles = doc.styles.element
    targets = [styles.find(f"{qn('w:docDefaults')}/{qn('w:rPrDefault')}/{qn('w:rPr')}")]
    targets += [s.get_or_add_rPr() for s in (st.element for st in doc.styles) if s.tag == qn("w:style")]
    for rpr in targets:
        if rpr is None:
            continue
        el = rpr.find(qn("w:lang"))
        if el is None:
            el = OxmlElement("w:lang")
            rpr.append(el)
        el.set(qn("w:val"), lang)
        el.set(qn("w:eastAsia"), lang)


# ---------- helpers ----------
def set_cell_borders(cell, **edges):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge, sz in edges.items():
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single" if sz else "nil")
        if sz:
            el.set(qn("w:sz"), str(sz))
            el.set(qn("w:color"), "000000")
        borders.append(el)


def shade(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def add_field(paragraph, instr):
    run = paragraph.add_run()
    for kind, text in (("begin", None), (None, instr), ("separate", None), (None, "1"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        elif text == instr:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        else:
            el = OxmlElement("w:t")
            el.text = text
        run._r.append(el)


def format_tables(doc):
    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl_pr = table._tbl.tblPr
        for old in tbl_pr.findall(qn("w:tblBorders")):
            tbl_pr.remove(old)
        width = OxmlElement("w:tblW")
        width.set(qn("w:w"), "5000")
        width.set(qn("w:type"), "pct")
        for old in tbl_pr.findall(qn("w:tblW")):
            tbl_pr.remove(old)
        tbl_pr.append(width)
        n_rows = len(table.rows)
        for r, row in enumerate(table.rows):
            if r == 0:
                tr_pr = row._tr.get_or_add_trPr()
                hdr = OxmlElement("w:tblHeader")
                tr_pr.append(hdr)
            for cell in row.cells:
                edges = {"left": 0, "right": 0, "top": 0, "bottom": 0}
                if r == 0:
                    edges.update(top=12, bottom=6)
                    shade(cell, TABLE_SHADE)
                if r == n_rows - 1:
                    edges["bottom"] = 12
                set_cell_borders(cell, **edges)
                for p in cell.paragraphs:
                    p.paragraph_format.line_spacing = 1.0
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.space_before = Pt(2)
                    for run in p.runs:
                        run.font.size = Pt(9.5)
                        if r == 0:
                            run.font.bold = True


def format_captions_and_figures(doc):
    paras = doc.paragraphs
    for i, p in enumerate(paras):
        text = p.text.strip()
        if re.match(r"^(Table|Figure) [A-Z]?\d+\s", text) or re.fullmatch(r"(Table|Figure) [A-Z]?\d+", text):
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
        if p._p.xpath(".//w:drawing"):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(8)
            if i > 0:
                paras[i - 1].paragraph_format.keep_with_next = True


def format_references(doc):
    in_refs = False
    for p in doc.paragraphs:
        if p.style.name.startswith("Heading") and p.text.strip() == "References":
            in_refs = True
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.page_break_before = True
            continue
        if p.style.name.startswith("Heading") and p.text.strip().startswith("Appendix"):
            in_refs = False
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.page_break_before = True
            continue
        if in_refs and p.text.strip():
            pf = p.paragraph_format
            pf.left_indent = Cm(1.27)
            pf.first_line_indent = Cm(-1.27)
            pf.line_spacing = 1.5
            pf.space_after = Pt(4)


def highlight_placeholders(doc):
    """Split runs so only the [[...]] text itself is highlighted."""
    def paragraphs():
        yield from doc.paragraphs
        for t in doc.tables:
            for row in t.rows:
                for cell in row.cells:
                    yield from cell.paragraphs

    for p in paragraphs():
        for run in list(p.runs):
            if "[[" not in run.text:
                continue
            parts = [s for s in re.split(r"(\[\[.*?\]\])", run.text) if s]
            run.text = parts[0]
            if parts[0].startswith("[["):
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW
            anchor = run._r
            for part in parts[1:]:
                new_r = copy.deepcopy(run._r)
                anchor.addnext(new_r)
                anchor = new_r
                new_run = Run(new_r, p)
                new_run.text = part
                new_run.font.highlight_color = WD_COLOR_INDEX.YELLOW if part.startswith("[[") else None


def fit_images(doc, width_cm=15.9):
    for shape in doc.inline_shapes:
        ratio = Cm(width_cm) / shape.width
        shape.width = Cm(width_cm)
        shape.height = int(shape.height * ratio)


def space_after_tables(doc):
    for tbl in doc.element.body.iter(qn("w:tbl")):
        nxt = tbl.getnext()
        if nxt is not None and nxt.tag == qn("w:p"):
            Paragraph(nxt, doc._body).paragraph_format.space_before = Pt(12)


def add_cover(doc, meta, wc):
    first = doc.paragraphs[0]

    def line(text, size=BODY_PT, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.CENTER, after=6, color=None):
        p = first.insert_paragraph_before()
        p.alignment = align
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(text)
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.name = FONT
        if color:
            run.font.color.rgb = color
        return p

    for _ in range(4):
        line("", after=12)
    line(meta.get("subtitle", ""), size=13, color=NAVY, after=18)
    line(meta.get("title", ""), size=20, bold=True, color=NAVY, after=36)
    line("Unit", size=10, bold=True, after=0)
    line(meta.get("unit", ""), after=14)
    line("Unit Coordinator", size=10, bold=True, after=0)
    line(meta.get("lecturer", ""), after=14)
    line("Team Members", size=10, bold=True, after=0)
    for member in meta.get("team", []):
        line(member, after=2)
    line("", after=10)
    line("Submission Date", size=10, bold=True, after=0)
    line(meta.get("date", ""), after=14)
    brk = line("", after=0)
    brk.runs[0].add_break(WD_BREAK.PAGE)


def add_page_numbers(doc, skip_first=True):
    section = doc.sections[0]
    section.different_first_page_header_footer = skip_first
    header = section.header
    p = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_field(p, "PAGE")


def set_page(doc):
    for s in doc.sections:
        s.page_width, s.page_height = Cm(21.0), Cm(29.7)
        s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Cm(2.54)


def build(md_path: str, out_path: str) -> None:
    md_file = Path(md_path).resolve()
    meta, body = split_front_matter(md_file.read_text(encoding="utf-8"))
    apply_theme(meta)
    wc = count(str(md_file))
    with tempfile.TemporaryDirectory() as tmp:
        ref = Path(tmp) / "reference.docx"
        make_reference_doc(ref)
        src = Path(tmp) / "body.md"
        src.write_text(prepare_markdown(body), encoding="utf-8")
        raw = Path(tmp) / "raw.docx"
        subprocess.run(
            ["pandoc", str(src), "-f", "markdown+autolink_bare_uris", "-o", str(raw),
             "--reference-doc", str(ref), "--resource-path", str(md_file.parent)],
            check=True,
        )
        doc = Document(str(raw))
    set_page(doc)
    format_tables(doc)
    space_after_tables(doc)
    fit_images(doc)
    format_captions_and_figures(doc)
    format_references(doc)
    has_cover = meta.get("cover_page", "true") != "false"  # front matter: cover_page: false
    if has_cover:
        add_cover(doc, meta, wc)
    highlight_placeholders(doc)
    add_page_numbers(doc, skip_first=has_cover)
    doc.core_properties.title = meta.get("title", "")
    doc.core_properties.author = "ExplainIDS team"
    doc.save(out_path)
    print(f"Wrote {out_path}  (body prose: {wc['prose']} words; tables: {wc['tables']})")


if __name__ == "__main__":
    build(sys.argv[1], sys.argv[2])
