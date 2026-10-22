"""Pre-submission checks for a report: placeholders, declarations, Word artefacts,
citation <-> reference consistency, word count and reference URLs.

Usage: python reports/tools/check_submission.py A1 [path/to/submission.docx] [--urls]
"""
import json
import re
import sys
import urllib.request
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).parent))
from wordcount import count  # noqa: E402

ROOT = Path(__file__).resolve().parents[2]
W14 = "http://schemas.microsoft.com/office/word/2010/wordml"
issues = []


def flag(msg):
    issues.append(msg)
    print("  ISSUE:", msg)


def all_text(doc):
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                parts.append(c.text)
    for s in doc.sections:
        parts += [p.text for p in s.header.paragraphs] + [p.text for p in s.footer.paragraphs]
    return "\n".join(parts)


def check_docx(path):
    print(f"\n== Word file: {path.name}")
    doc = Document(str(path))
    body = doc.element.body
    text = all_text(doc)

    for pat in (r"\[\[", r"PLACEHOLDER", r"\[insert", r"\[describe", r"TODO", r"TBC", r"XX"):
        n = len(re.findall(pat, text))
        if n:
            flag(f"{n} occurrence(s) of '{pat}' left in the document")

    hl = [r.text for p in body.iter(qn("w:p")) for r in p.iter(qn("w:r"))
          if r.find(f"{qn('w:rPr')}/{qn('w:highlight')}") is not None]
    if hl:
        flag(f"highlighted text still present: {[''.join(t.text or '' for t in r.iter(qn('w:t'))) for r in []] or len(hl)} run(s)")

    labels = ["Student Declaration", "Reference Declaration", "AI: NOT included", "AI: included + acknowledged"]
    boxes = list(body.iter(qn("w:sdt")))
    states = {}
    for label, box in zip(labels, boxes):
        c = box.find(f".//{{{W14}}}checked")
        states[label] = c is not None and c.get(f"{{{W14}}}val") in ("1", "true")
        print(f"  checkbox {label:30s}: {'TICKED' if states[label] else 'not ticked'}")
    if boxes:
        if not states.get("Student Declaration"):
            flag("Student Declaration is not ticked")
        if not states.get("Reference Declaration"):
            flag("Reference Declaration is not ticked")
        if states.get("AI: NOT included") == states.get("AI: included + acknowledged"):
            flag("AI Declaration: exactly one option must be ticked")
        if states.get("AI: NOT included"):
            flag("AI Declaration says NO AI content, but AI tools were used (see acknowledgements)")

    xml = body.xml
    for tag, name in (("<w:ins ", "tracked insertions"), ("<w:del ", "tracked deletions"),
                      ("commentReference", "comments"), ("<w:vanish", "hidden text")):
        if xml.count(tag):
            flag(f"{xml.count(tag)} {name}")

    sizes = [(round(s.page_width.cm, 1), round(s.page_height.cm, 1)) for s in doc.sections]
    print(f"  sections/page sizes: {sizes}; images: {len(doc.inline_shapes)}")
    cp = doc.core_properties
    print(f"  properties: title={cp.title[:50]!r} author={cp.author!r} last_modified_by={cp.last_modified_by!r}")
    return text


def check_markdown(md_path, word_range=(1350, 1650)):
    print(f"\n== Source: {md_path.relative_to(ROOT)}")
    raw = md_path.read_text(encoding="utf-8")
    body_md, _, refs_md = raw.partition("\n# References")
    refs_md = refs_md.split("\n# Appendix")[0]
    body_md = re.sub(r"^---\n.*?\n---\n", "", body_md, flags=re.S)
    if "[[" in raw:
        flag("placeholders '[[' remain in the Markdown source")

    wc = count(str(md_path))
    print(f"  words: prose {wc['prose']} | tables {wc['tables']} | captions {wc['captions']} | prose+tables {wc['prose'] + wc['tables']}")
    lo, hi = word_range
    if not lo <= wc["prose"] <= hi:
        flag(f"body prose {wc['prose']} is outside {lo}-{hi}")

    refs = [r.strip() for r in refs_md.strip().split("\n\n") if r.strip()]
    keys = []
    for r in refs:
        m = re.match(r"^\*?([^(*]+?)\*?\s*(?:\((\d{4}|n\.d\.)\)|\(Cth\))", r)
        if m:
            year = m.group(2) or re.search(r"(\d{4})", m.group(1)).group(1)
            keys.append((m.group(1).split(",")[0].strip().rstrip("."), year, r[:70]))
        else:
            flag(f"could not parse reference: {r[:70]}")
    sorted_refs = sorted(refs, key=lambda s: re.sub(r"[^a-z0-9 ]", "", s.lower()))  # APA: letter by letter
    if refs != sorted_refs:
        flag("reference list is not in alphabetical order")

    cited = set()
    for m in re.finditer(r"\(([^()]*?\d{4}[a-z]?|[^()]*?n\.d\.)\)", body_md):
        for part in m.group(1).split(";"):
            cited.add(part.strip())
    narrative = re.findall(r"([A-Z][A-Za-z\-]+(?: (?:et al\.|and|&) ?[A-Z]?[A-Za-z\-]*)?)(?:'s)? \((\d{4})\)", body_md)

    def surname_in_body(s, y):
        s0 = s.split()[0]
        abbreviations = {"Australian Signals Directorate": "ASD"}
        names = [s0, abbreviations.get(s, s0)]
        if s.startswith("*"):
            names = [s.strip("*")]
        return any(re.search(re.escape(n) + r".{0,80}" + re.escape(y), body_md, re.S) for n in names)

    for surname, year, short in keys:
        if surname.startswith("*") or "Act" in surname:
            name = surname.strip("*")
            if name.split(" (")[0] not in body_md:
                flag(f"legislation not cited in text: {name}")
            continue
        if not surname_in_body(surname, year):
            flag(f"reference apparently not cited in text: {short}")

    ref_names = " ".join(k[0] for k in keys) + " ASD"
    for c in sorted(cited):
        m = re.match(r"(?:e\.g\., )?([A-Za-z][\w\-\s&.]*?),? (\d{4}|n\.d\.)", c)
        if not m:
            continue
        first = m.group(1).split()[0]
        if first not in ref_names:
            flag(f"in-text citation without reference entry: ({c})")
    print(f"  references: {len(refs)} | parenthetical citation groups checked: {len(cited)} | narrative: {len(narrative)}")
    return refs


def check_urls(refs):
    print("\n== Reference URLs")
    for r in refs:
        for url in re.findall(r"https?://\S+", r):
            url = url.rstrip(".")
            req = urllib.request.Request(url, method="GET", headers={"User-Agent": "Mozilla/5.0 (reference check)"})
            try:
                with urllib.request.urlopen(req, timeout=25) as resp:
                    status = resp.status
            except Exception as e:  # noqa: BLE001
                status = getattr(e, "code", type(e).__name__)
            ok = status in (200, 301, 302, 303) or (status == 202 and "doi.org/10.1109" in url)  # IEEE bot-check returns 202
            print(f"  {'OK ' if ok else '?? '} {status}  {url}")
            if not ok:
                flag(f"URL did not return 200 (may be bot-blocking; check in a browser): {url} -> {status}")


def main():
    assessment = sys.argv[1]
    folder = ROOT / "reports" / assessment
    cfg = json.loads((folder / "cover_sheet.json").read_text(encoding="utf-8"))
    refs = check_markdown(folder / cfg["report_markdown"], tuple(cfg.get("word_range", (1350, 1650))))
    docx_args = [a for a in sys.argv[2:] if not a.startswith("--")]
    targets = [Path(a) for a in docx_args] or sorted((folder / "submission").glob("*.docx"))
    for t in targets:
        check_docx(t)
    if "--urls" in sys.argv:
        check_urls(refs)
    print(f"\n{len(issues)} issue(s) found.")


if __name__ == "__main__":
    main()
